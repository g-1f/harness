from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'outputs/library-harness-design.md'
TARGET = ROOT / 'outputs/library-harness-design.docx'


def diagram():
    path = ROOT / 'work/architecture.png'
    im = Image.new('RGB', (1800, 620), 'white')
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/calibri.ttf', 31)
    small = ImageFont.truetype('C:/Windows/Fonts/calibri.ttf', 25)
    boxes = [(30, 55, 340, 180, 'Authenticated\nrequest'),
             (460, 55, 820, 180, 'Supervisor\nand shared ledger'),
             (950, 55, 1320, 180, 'Deep Agents worker\nand QuickJS'),
             (1430, 55, 1770, 180, 'Skill graph\nand evidence'),
             (950, 380, 1320, 515, 'Candidate\nand required reviews'),
             (1430, 380, 1770, 515, 'Accepted artifact\nor needs review')]
    for x1, y1, x2, y2, label in boxes:
        draw.rectangle((x1, y1, x2, y2), fill='#f3f5f7', outline='#6d7982', width=2)
        draw.multiline_text(((x1+x2)/2, (y1+y2)/2), label, font=font,
                            fill='black', anchor='mm', align='center', spacing=5)

    def arrow(points):
        draw.line(points, fill='#3c4852', width=4)
        (x0, y0), (x, y) = points[-2:]
        if x > x0:
            tri = [(x, y), (x-14, y-7), (x-14, y+7)]
        elif x < x0:
            tri = [(x, y), (x+14, y-7), (x+14, y+7)]
        elif y > y0:
            tri = [(x, y), (x-7, y-14), (x+7, y-14)]
        else:
            tri = [(x, y), (x-7, y+14), (x+7, y+14)]
        draw.polygon(tri, fill='#3c4852')

    arrow([(340, 118), (460, 118)])
    arrow([(820, 118), (950, 118)])
    arrow([(1320, 118), (1430, 118)])
    arrow([(1135, 180), (1135, 380)])
    arrow([(1320, 450), (1430, 450)])
    arrow([(1000, 180), (1000, 270), (640, 270), (640, 180)])
    arrow([(950, 450), (640, 450), (640, 290)])
    draw.text((765, 240), 'Recursive calls', font=small, fill='#3c4852', anchor='mm')
    draw.text((790, 480), 'Critic calls', font=small, fill='#3c4852', anchor='mm')
    draw.text((1255, 285), 'Submit', font=small, fill='#3c4852', anchor='mm')
    im.save(path)
    return path


doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.68)
sec.left_margin = sec.right_margin = Inches(.72)
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.footer_distance = Inches(.3)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)

normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.widow_control = True
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True

if 'Code' not in doc.styles:
    from docx.enum.style import WD_STYLE_TYPE
    doc.styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
code_style = doc.styles['Code']
code_style.font.name = 'Consolas'
code_style.font.size = Pt(8.5)
code_style.font.color.rgb = RGBColor(0, 0, 0)
code_style.paragraph_format.space_after = Pt(7)
code_style.paragraph_format.line_spacing = 1.0


def inline(paragraph, text):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)'
    for token in re.split(pattern, text):
        link = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', token)
        if link:
            relation = paragraph.part.relate_to(link[2], RT.HYPERLINK, is_external=True)
            node = OxmlElement('w:hyperlink')
            node.set(qn('r:id'), relation)
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '28506C')
            props.append(color)
            underline = OxmlElement('w:u')
            underline.set(qn('w:val'), 'single')
            props.append(underline)
            run.append(props)
            t = OxmlElement('w:t')
            t.text = link[1]
            run.append(t)
            node.append(run)
            paragraph._p.append(node)
        elif token.startswith('`'):
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
        elif token.startswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token)


def make_table(lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r'[-: ]+', c) for c in r)]
    cols = len(rows[0])
    table = doc.add_table(rows=0, cols=cols)
    table.autofit = False
    widths = [1.6, 2.42, 3.04] if cols == 3 else [2.0, 5.06]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, values in enumerate(rows):
        row = table.add_row()
        for j, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + side)
                item.set(qn('w:val'), 'single')
                item.set(qn('w:sz'), '4')
                item.set(qn('w:color'), 'D9D9D9')
                borders.append(item)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + side)
                item.set(qn('w:w'), '90')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            props.append(margins)
            if i == 0 or i % 2 == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'DFE6EB' if i == 0 else 'F5F7F8')
                props.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.03
            inline(p, text)
            for r in p.runs:
                r.font.size = Pt(9)
                r.bold = i == 0
        pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        pr.append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
diagram_path = diagram()
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('```'):
        code = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            code.append(lines[i])
            i += 1
        p = doc.add_paragraph('\n'.join(code), 'Code')
        p.paragraph_format.keep_together = True
    elif line.startswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].startswith('|'):
            table_lines.append(lines[i])
            i += 1
        make_table(table_lines)
        continue
    else:
        style = 'Normal'
        if i == 2:
            style = 'Subtitle'
        elif re.match(r'^\d+\. ', line):
            style = 'List Number'
            line = re.sub(r'^\d+\. ', '', line)
        elif line.startswith('- '):
            style = 'List Bullet'
            line = line[2:]
        p = doc.add_paragraph(style=style)
        inline(p, line)
        if line.startswith('The supplied `architecture.mmd`'):
            pic = doc.add_paragraph()
            pic.add_run().add_picture(str(diagram_path), width=Inches(7.0))
            pic.paragraph_format.keep_with_next = True
            cap = doc.add_paragraph('Figure 1  Supervised recursive execution and result review')
            cap.paragraph_format.space_after = Pt(8)
            for r in cap.runs:
                r.italic = True
                r.font.size = Pt(9)
    i += 1

p = sec.footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
p._p.append(field)
doc.core_properties.title = 'The Library Harness Design'
doc.core_properties.subject = 'Skill graph harness using Deep Agents and programmatic recursive execution'
doc.core_properties.author = 'Library design review'
doc.save(TARGET)
print(TARGET)
print(f'{len(doc.paragraphs)} paragraphs; {len(doc.tables)} tables')
